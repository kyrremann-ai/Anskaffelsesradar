"""Kandidatlister fra Valgdirektoratet (valg.no, «Valglister og kandidater», CSV/XLSX).

Hentes automatisk ved hver oppdatering (refresh): kommunestyre- og fylkestingsvalget 2023 og
stortingsvalget 2025, med «innvalgt» satt fra Valgdirektoratets lister over valgte.
Kan også importeres fra nedlastet fil:
    python -m radar.ingest.valg kandidater-2023.csv --valg "Kommunestyrevalget 2023"

Overskriftene matches løst: navn, fødselsår, kommunenummer/kommune.
Personvern: vi lagrer bare normalisert navn, fødselsår, kommune, valgt/kandidat og valg. Parti leses ikke inn
(politisk oppfatning er en særlig kategori etter GDPR art. 9). Treff vises som
«mulig kobling» og må alltid verifiseres manuelt (navnelikhet). Gjør en DPIA før produksjon.
"""
from __future__ import annotations

import argparse
import csv
import io
import re
from pathlib import Path

from ..db import connect, norm_name, tx


def _find(h: list[str], *keys) -> int | None:
    # Eksakt treff først, ellers delstreng («navn» skal ikke treffe «partinavn» når «navn» finnes)
    for k in keys:
        if k in h:
            return h.index(k)
    for k in keys:
        for i, c in enumerate(h):
            if k in c:
                return i
    return None


VALG_URL = "https://www.valg.no/globalassets"
# (valg, kandidatliste, liste over valgte). Valgte-filene gir «innvalgt».
SOURCES = [
    ("Kommunestyrevalget 2023",
     "/dokumenter-2023/om-valg/valgdata/lister-og-kandidater/kommunestyrevalget_2023_lister_og_kandidater_270623.xlsx",
     "/dokumenter-2023/om-valg/valgdata/valgte-kandidater/valgte_kandidater_kommunestyre_2023---med-forbehold-om-endringer-3.xlsx"),
    ("Fylkestingsvalget 2023",
     "/dokumenter-2023/om-valg/valgdata/lister-og-kandidater/fylkestingsvalget_2023_lister_og_kandidater_270623.xlsx",
     "/dokumenter-2023/om-valg/valgdata/valgte-kandidater/valgte_kandidater_fylkesting_2023---med-forbehold-om-endringer.xlsx"),
    ("Stortingsvalget 2025",
     "/valgdata/lister-og-kandidater/valgdistrikt/lister_og_kandidater_stortingsvalget_2025.csv",
     "/dokumenter-2025/om-valg/valgdata/valgte_kandidater_stortingsvalget_2025---komplett.xlsx"),
]


def _cell(v) -> str:
    if v is None:
        return ""
    return v.isoformat() if hasattr(v, "isoformat") else str(v)


def read_rows(raw: bytes, name: str) -> list[list[str]]:
    """Rader fra CSV eller XLSX som lister av tekst."""
    if name.lower().endswith(".xlsx"):
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(raw), read_only=True).worksheets[0]
        return [[_cell(v) for v in r] for r in ws.iter_rows(values_only=True)]
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("latin-1")
    delim = ";" if text.count(";") > text.count(",") else ","
    return list(csv.reader(io.StringIO(text), delimiter=delim))


def parse(text: str) -> list[dict]:
    delim = ";" if text.count(";") > text.count(",") else ","
    return parse_rows(list(csv.reader(io.StringIO(text), delimiter=delim)))


def parse_rows(rows: list[list[str]]) -> list[dict]:
    h = [c.strip().lower() for c in rows[0]]
    i_navn = _find(h, "kandidatnavn", "navn", "kandidat")
    i_for = _find(h, "fornavn")
    i_etter = _find(h, "etternavn")
    i_aar = _find(h, "fødselsår", "fodselsar", "fødselsdato", "født", "fodt")
    i_knr = _find(h, "kommunenummer", "kommunenr")
    i_kom = _find(h, "kommunenavn", "kommune")
    i_valgt = _find(h, "valgt", "innvalgt")
    out = []
    for r in rows[1:]:
        if not r:
            continue
        def g(i, r=r):
            return r[i].strip() if i is not None and i < len(r) else ""
        navn = f"{g(i_for)} {g(i_etter)}".strip() if i_for is not None and i_etter is not None else g(i_navn)
        y = re.search(r"(?:19|20)\d{2}", g(i_aar))
        if not navn or not y:
            continue
        aar = y.group(0)
        knr = re.sub(r"\D", "", g(i_knr)).zfill(4) if g(i_knr) else None
        out.append({"navn_norm": norm_name(navn), "fodselsaar": int(aar), "kommunenr": knr, "kommune": g(i_kom),
                    "valgt": int(g(i_valgt).lower() in ("ja", "1", "true", "x", "valgt"))})
    return out


def _elected(rows: list[list[str]]) -> tuple[set, set]:
    """Nøkler for innvalgte: (navn, år, kommunenr) når fødselsdato finnes, ellers bare navn."""
    h = [c.strip().lower() for c in rows[0]]
    i_navn, i_dato = _find(h, "kandidatnavn", "navn"), _find(h, "fødselsdato")
    i_knr, i_inn = _find(h, "kommunenummer"), _find(h, "innvalgt")
    full, names = set(), set()
    for r in rows[1:]:
        if not r or (i_inn is not None and str(r[i_inn]).strip().lower() != "ja"):
            continue
        n = norm_name(r[i_navn])
        y = re.search(r"(?:19|20)\d{2}", r[i_dato]) if i_dato is not None else None
        if y:
            full.add((n, int(y.group(0)), r[i_knr].zfill(4) if i_knr is not None and r[i_knr] else None))
        else:
            names.add(n)
    return full, names


def refresh(con, log=print) -> int:
    """Laster ned kandidatlistene fra valg.no og lagrer dem. Returnerer antall kandidater."""
    import httpx
    total = 0
    with httpx.Client(headers={"User-Agent": "Anskaffelsesradar/2.0"}, follow_redirects=True, timeout=120) as client:
        for valg, lists, elected in SOURCES:
            def get(path):
                r = client.get(VALG_URL + path)
                r.raise_for_status()
                return read_rows(r.content, path)
            rows = parse_rows(get(lists))
            full, names = _elected(get(elected))
            for r in rows:
                r["valgt"] = int((r["navn_norm"], r["fodselsaar"], r["kommunenr"]) in full or r["navn_norm"] in names)
            store(con, rows, valg)
            log(f"  {valg}: {len(rows)} kandidater, {sum(r['valgt'] for r in rows)} innvalgt")
            total += len(rows)
    return total


def store(con, rows: list[dict], valg: str) -> None:
    with tx(con):
        con.execute("DELETE FROM politicians WHERE valg=?", (valg,))
        con.executemany("INSERT INTO politicians(navn_norm, fodselsaar, kommunenr, kommune, valgt, valg) VALUES (?,?,?,?,?,?)",
                        [(r["navn_norm"], r["fodselsaar"], r["kommunenr"], r["kommune"], r["valgt"], valg) for r in rows])


if __name__ == "__main__":
    from ..config import settings
    ap = argparse.ArgumentParser()
    ap.add_argument("csv", type=Path)
    ap.add_argument("--valg", default="Kommunestyrevalget 2023")
    a = ap.parse_args()
    rows = parse_rows(read_rows(a.csv.read_bytes(), a.csv.name))
    store(connect(settings.db_path), rows, a.valg)
    print(f"Importerte {len(rows)} kandidater")
